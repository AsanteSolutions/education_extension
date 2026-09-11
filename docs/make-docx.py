#!/usr/bin/env python3
"""Turn the markdown in this folder into Word documents.

The markdown is the source. These are copies for people who would rather be
handed a file than sent a link, so they are regenerated rather than edited:
anything typed into the .docx is lost the next time this runs.

    pip install python-docx
    python3 docs/make-docx.py

Handles the subset of markdown these documents actually use -- headings,
paragraphs, bullet and numbered lists, pipe tables, block quotes, fenced code,
rules, and inline bold/italic/code/links.
"""

import os
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is not installed. Run: pip install python-docx")


HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "word")

# The institution's green, the same one on the Proof of Registration.
ACCENT = RGBColor(0x1A, 0x7A, 0x3C)
INK = RGBColor(0x1C, 0x23, 0x21)
SOFT = RGBColor(0x4A, 0x56, 0x52)
WARN = RGBColor(0xB4, 0x53, 0x09)

# Fonts that exist on the reader's machine rather than ones we would have to
# ship: a serif for headings and a humanist sans for text, as in the web copy.
HEADING_FONT = "Cambria"
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

# The order they belong in, which is not alphabetical.
ORDER = [
    "README.md",
    "for-students.md",
    "for-registrars.md",
    "for-lecturers-and-qa.md",
    "settings.md",
    "for-administrators.md",
    "glossary.md",
]

TITLES = {
    "README.md": "Education Extension",
    "for-students.md": "Education Extension — For Students",
    "for-registrars.md": "Education Extension — For Registrars",
    "for-lecturers-and-qa.md": "Education Extension — For Lecturers and QA",
    "settings.md": "Education Extension — Settings Reference",
    "for-administrators.md": "Education Extension — For Administrators",
    "glossary.md": "Education Extension — Glossary",
}


# --------------------------------------------------------------------------
# document setup
# --------------------------------------------------------------------------


def shade(cell, colour):
    """Fill a table cell. python-docx has no API for it, so set the XML."""
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:fill"), colour)
    cell._tc.get_or_add_tcPr().append(fill)


def rule_below(paragraph, colour="D6DED8"):
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), colour)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def new_document(title):
    doc = Document()

    for section in doc.sections:
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.1)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

        footer = section.footer.paragraphs[0]
        footer.text = "Education Extension  ·  Tsolo Agriculture & Rural Development Institute"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.name = BODY_FONT
            run.font.size = Pt(8)
            run.font.color.rgb = SOFT

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    sizes = {"Heading 1": 20, "Heading 2": 15, "Heading 3": 12.5, "Heading 4": 11}
    for name, size in sizes.items():
        style = doc.styles[name]
        style.font.name = HEADING_FONT if name != "Heading 4" else BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT if name in ("Heading 1", "Heading 2") else INK
        style.paragraph_format.space_before = Pt(16 if name == "Heading 2" else 12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    heading = doc.add_paragraph()
    run = heading.add_run(title)
    run.font.name = HEADING_FONT
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = ACCENT
    heading.paragraph_format.space_after = Pt(10)
    rule_below(heading)

    return doc


# --------------------------------------------------------------------------
# inline formatting
# --------------------------------------------------------------------------

INLINE = re.compile(
    r"(\*\*.+?\*\*)"        # bold
    r"|(`[^`]+`)"           # code
    r"|(\*[^*]+\*)"         # italic
    r"|(\[[^\]]+\]\([^)]+\))"  # link
)


def write_inline(paragraph, text, base_bold=False):
    """Add `text` to `paragraph`, honouring the inline markdown in it."""
    for piece in INLINE.split(text):
        if not piece:
            continue

        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(9.5)
            run.font.color.rgb = SOFT
        elif piece.startswith("*") and piece.endswith("*"):
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
        elif piece.startswith("["):
            label, _, target = piece[1:].partition("](")
            target = target.rstrip(")")
            run = paragraph.add_run(label)
            run.font.color.rgb = ACCENT
            # An anchor into a sibling markdown file means nothing in Word, so
            # only a real address is worth showing.
            if target.startswith("http"):
                trailing = paragraph.add_run(" ({0})".format(target))
                trailing.font.size = Pt(9)
                trailing.font.color.rgb = SOFT
        else:
            run = paragraph.add_run(piece)

        if base_bold:
            run.bold = True


# --------------------------------------------------------------------------
# block parsing
# --------------------------------------------------------------------------


def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True

    for cell, text in zip(table.rows[0].cells, header):
        cell.paragraphs[0].text = ""
        write_inline(cell.paragraphs[0], text, base_bold=True)
        cell.paragraphs[0].paragraph_format.space_after = Pt(3)
        shade(cell, "EAF3EC")

    for line in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, line):
            cell.paragraphs[0].text = ""
            write_inline(cell.paragraphs[0], text)
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_quote(doc, lines):
    paragraph = doc.add_paragraph()
    write_inline(paragraph, " ".join(lines))
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)

    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "B45309")
    border.append(left)
    paragraph._p.get_or_add_pPr().append(border)

    for run in paragraph.runs:
        if run.font.color.rgb is None:
            run.font.color.rgb = SOFT


def add_code(doc, lines):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("\n".join(lines))
    run.font.name = MONO_FONT
    run.font.size = Pt(9)
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.0


def convert(path, doc):
    with open(path, encoding="utf-8") as handle:
        lines = handle.read().splitlines()

    index = 0
    paragraph_buffer = []

    def flush():
        if not paragraph_buffer:
            return
        paragraph = doc.add_paragraph()
        write_inline(paragraph, " ".join(paragraph_buffer))
        paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            flush()
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code(doc, block)
            index += 1
            continue

        # table
        if stripped.startswith("|") and index + 1 < len(lines) and set(
            lines[index + 1].strip()
        ) <= set("|-: "):
            flush()
            rows = [split_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            add_table(doc, rows)
            continue

        # block quote
        if stripped.startswith(">"):
            flush()
            block = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                block.append(lines[index].strip().lstrip(">").strip())
                index += 1
            add_quote(doc, [text for text in block if text])
            continue

        # heading
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            flush()
            level = len(heading.group(1))
            # The file's own H1 is already the document title.
            if not (level == 1 and doc.paragraphs and len(doc.paragraphs) <= 1):
                paragraph = doc.add_heading(level=min(level, 4))
                paragraph.text = ""
                write_inline(paragraph, heading.group(2))
            index += 1
            continue

        # rule
        if stripped in ("---", "***", "___"):
            flush()
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(2)
            rule_below(spacer)
            index += 1
            continue

        # lists
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        number = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if bullet or number:
            flush()
            style = "List Number" if number else "List Bullet"
            paragraph = doc.add_paragraph(style=style)
            write_inline(paragraph, (number or bullet).group(1))
            paragraph.paragraph_format.space_after = Pt(4)
            index += 1
            continue

        if not stripped:
            flush()
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush()


def main():
    if not os.path.isdir(OUT):
        os.makedirs(OUT)

    names = [name for name in ORDER if os.path.exists(os.path.join(HERE, name))]
    names += sorted(
        name
        for name in os.listdir(HERE)
        if name.endswith(".md") and name not in ORDER
    )

    for name in names:
        title = TITLES.get(name, name[:-3].replace("-", " ").title())
        doc = new_document(title)
        convert(os.path.join(HERE, name), doc)

        out = os.path.join(OUT, name[:-3] + ".docx")
        doc.save(out)
        print("  {0:<34} -> {1}".format(name, os.path.relpath(out, HERE)))

    print("\n{0} documents written to {1}/".format(len(names), os.path.relpath(OUT, HERE)))


if __name__ == "__main__":
    main()
